"""
Company Profile Loader
Reads Chipscape's own turnover, credentials, past-project data and
subcontractor/system-integrator partner credentials from the database
(managed via the dashboard's Admin section — see database/db.py) and turns
them into a compact summary that can be handed to the AI for
tender-eligibility matching (see eligibility_matcher.py).

The loaders below (_load_turnover etc.) also still read the original Excel
files in company_data/ — kept only so scripts/migrate_company_data_to_db.py
has a one-time source to migrate from. Nothing at runtime reads Excel
anymore; the database is the source of truth.
"""

import os
from datetime import datetime

COMPANY_DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "company_data")
TURNOVER_FILE = os.path.join(COMPANY_DATA_DIR, "Chipscape Turnover.xlsx")
REGISTER_FILE = os.path.join(COMPANY_DATA_DIR, "Eligibility_Master_Register.xlsx")

_cached_profile = None


def _fmt_amount(value):
    if value is None:
        return ""
    try:
        return f"Rs {float(value):,.0f}"
    except (TypeError, ValueError):
        return str(value)


def _fmt_date(value):
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y")
    return str(value) if value else ""


def _load_turnover():
    """Financial-year turnover history, most recent year last. Only used by
    the one-time Excel migration script, not at runtime."""
    import openpyxl
    years = []
    if not os.path.exists(TURNOVER_FILE):
        return years
    wb = openpyxl.load_workbook(TURNOVER_FILE, data_only=True)
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        fy = row[0]
        if not fy or not isinstance(fy, str) or "-" not in fy:
            continue
        turnover, gst, total, status = row[1], row[2], row[3], row[4]
        if turnover is None:
            continue
        years.append({
            "financial_year": fy,
            "turnover": float(turnover),
            "total_with_gst": float(total) if total else float(turnover),
            "status": status or "",
        })
    return years


def _load_credentials():
    """Registrations, licenses and certifications with validity. Only used
    by the one-time Excel migration script, not at runtime."""
    import openpyxl
    creds = []
    if not os.path.exists(REGISTER_FILE):
        return creds
    wb = openpyxl.load_workbook(REGISTER_FILE, data_only=True)
    if "Tab 2 Company Credentials" not in wb.sheetnames:
        return creds
    ws = wb["Tab 2 Company Credentials"]
    today = datetime.now()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            break
        name, ctype, issuer, cert_no, issue_date, expiry_date, status = (
            row[2], row[3], row[4], row[5], row[6], row[7], row[8]
        )
        expired = bool(isinstance(expiry_date, datetime) and expiry_date < today)
        creds.append({
            "name": name or "",
            "type": ctype or "",
            "issuing_authority": issuer or "",
            "certificate_no": str(cert_no) if cert_no is not None else "",
            "issue_date": _fmt_date(issue_date),
            "expiry_date": _fmt_date(expiry_date) if expiry_date else "No expiry",
            "status": "Expired" if expired else (status or "Valid"),
        })
    return creds


def _load_project_references():
    """Past completed projects with category, tags, and value — used to
    judge relevant prior experience against a new tender's scope. Only used
    by the one-time Excel migration script, not at runtime."""
    import openpyxl
    projects = []
    if not os.path.exists(REGISTER_FILE):
        return projects
    wb = openpyxl.load_workbook(REGISTER_FILE, data_only=True)
    if "Tab 1 Project References" not in wb.sheetnames:
        return projects
    ws = wb["Tab 1 Project References"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            break
        projects.append({
            "project_name": row[2] or "",
            "customer": row[3] or "",
            "department_type": row[4] or "",
            "location": row[6] or "",
            "solution_category": row[7] or "",
            "eligibility_tags": row[8] or "",
            "project_scope": row[9] or "",
            "project_value": float(row[10]) if row[10] else None,
            "government_project": row[12] or "",
            "financial_year": row[13] or "",
            "project_status": row[20] or "",
        })
    return projects


def _load_project_quantities():
    """Itemized quantities actually supplied on past projects (e.g. '20
    ANPR Cameras', '500 Wireless Nodes') — lets the matcher back up a
    quantity-based eligibility clause ("must have supplied 15+ units of X")
    with a real, specific past number instead of just a project name. Only
    used by the one-time Excel migration script, not at runtime."""
    import openpyxl
    quantities = []
    if not os.path.exists(REGISTER_FILE):
        return quantities
    wb = openpyxl.load_workbook(REGISTER_FILE, data_only=True)
    if "Tab 5 Project Quantity Register" not in wb.sheetnames:
        return quantities
    ws = wb["Tab 5 Project Quantity Register"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            break
        quantities.append({
            "project_name": row[3] or "",
            "solution_category": row[4] or "",
            "quantity_category": row[5] or "",
            "quantity_executed": row[6],
            "unit": row[7] or "",
            "customer": row[8] or "",
        })
    return quantities


def load_company_profile(force_reload: bool = False) -> dict:
    """Loads and caches Chipscape's turnover, credentials, project history
    and partner credentials from the database, ready to hand to the AI
    matcher."""
    global _cached_profile
    if _cached_profile is not None and not force_reload:
        return _cached_profile

    from database.db import (
        get_company_turnover_list, get_company_credentials_list,
        get_project_references_list, get_project_quantities_list,
        get_partners_list, get_partner_credentials_list,
    )

    turnover = get_company_turnover_list()
    credentials = get_company_credentials_list()
    projects = get_project_references_list()
    project_quantities = get_project_quantities_list()

    partners = get_partners_list()
    partner_credentials = get_partner_credentials_list()
    creds_by_partner = {}
    for pc in partner_credentials:
        creds_by_partner.setdefault(pc["partner_id"], []).append(pc)
    partners_with_credentials = [
        {**p, "credentials": creds_by_partner.get(p["id"], [])}
        for p in partners
    ]

    today = datetime.now()

    def _status_with_expiry_check(cred):
        expiry = cred.get("expiry_date")
        if expiry and expiry != "No expiry":
            try:
                if datetime.strptime(expiry, "%d-%m-%Y") < today:
                    return "Expired"
            except ValueError:
                pass
        return cred.get("status") or "Valid"

    for c in credentials:
        c["status"] = _status_with_expiry_check(c)
    for p in partners_with_credentials:
        for c in p["credentials"]:
            c["status"] = _status_with_expiry_check(c)

    audited_years = [y for y in turnover if (y.get("status") or "").upper() == "AUDITED"]

    _cached_profile = {
        "turnover_history": turnover,
        "latest_audited_turnover": audited_years[-1] if audited_years else None,
        "credentials": credentials,
        "project_references": projects,
        "project_quantities": project_quantities,
        "partners": partners_with_credentials,
    }
    return _cached_profile


def company_profile_as_text(profile: dict = None) -> str:
    """Compact plain-text rendering of the company profile for use inside
    an AI prompt — keeps token usage reasonable while including everything
    the matcher needs to reason about eligibility."""
    profile = profile or load_company_profile()

    lines = ["=== CHIPSCAPE TURNOVER HISTORY (Audited unless noted) ==="]
    for y in profile["turnover_history"]:
        lines.append(
            f"- {y['financial_year']}: {_fmt_amount(y['turnover'])} "
            f"(incl. GST: {_fmt_amount(y['total_with_gst'])}) [{y['status']}]"
        )

    # Cross-checked against the real certificate PDFs where possible (see
    # credential_verifier.py) — the AI is told explicitly whether each
    # credential was independently confirmed, so it can weigh a
    # verified certificate more confidently than one only recorded in the
    # register, and flag a genuine mismatch as a real risk rather than
    # missing it silently.
    try:
        from analysis.credential_verifier import verification_by_name
        verification = verification_by_name()
    except Exception as e:
        print(f"[Company Profile] Credential verification unavailable: {e}")
        verification = {}

    lines.append("\n=== CHIPSCAPE CERTIFICATIONS, LICENSES & REGISTRATIONS ===")
    for c in profile["credentials"]:
        v = verification.get(c["name"])
        v_note = f" | Cross-check: {v['status']}" if v else ""
        lines.append(
            f"- {c['name']} ({c['type']}) | Issuer: {c['issuing_authority']} | "
            f"Valid until: {c['expiry_date']} | Status: {c['status']}{v_note}"
        )

    if profile.get("partners"):
        lines.append(
            "\n=== SUBCONTRACTOR / SYSTEM-INTEGRATOR PARTNER CREDENTIALS ===\n"
            "(Only relevant if the tender allows bidding through a partner/OEM "
            "dealer relationship rather than requiring Chipscape's own credential.)"
        )
        for p in profile["partners"]:
            lines.append(f"- Partner: {p['name']} ({p.get('role') or 'Partner'})")
            for c in p["credentials"]:
                lines.append(
                    f"    - {c['name']} ({c['type']}) | Issuer: {c['issuing_authority']} | "
                    f"Valid until: {c['expiry_date']} | Status: {c['status']}"
                )

    lines.append("\n=== CHIPSCAPE PAST PROJECT EXPERIENCE (completed/ongoing) ===")
    for p in profile["project_references"]:
        lines.append(
            f"- {p['project_name']} | Customer: {p['customer']} ({p['department_type']}) | "
            f"Category: {p['solution_category']} | Tags: {p['eligibility_tags']} | "
            f"Value: {_fmt_amount(p['project_value'])} | Govt project: {p['government_project']} | "
            f"Status: {p['project_status']}"
        )

    if profile.get("project_quantities"):
        lines.append(
            "\n=== ITEMIZED QUANTITIES ACTUALLY SUPPLIED ON PAST PROJECTS ===\n"
            "(Use this to back up a quantity-based tender requirement, e.g. "
            "'must have supplied 15+ units of X', with a real specific number "
            "instead of just naming the project.)"
        )
        for q in profile["project_quantities"]:
            qty = q["quantity_executed"]
            try:
                qty_text = f"{float(qty):g}"
            except (TypeError, ValueError):
                qty_text = str(qty or "")
            lines.append(
                f"- {q['project_name']} | Category: {q['solution_category']} | "
                f"{q['quantity_category']}: {qty_text} {q['unit']} | Customer: {q['customer']}"
            )

    return "\n".join(lines)
